from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt


SOURCE = Path("analysis_outputs/dissertation_revisions/Messinger_Dissertation_Revised_WildBird_NullFindings.docx")
OUT = Path("analysis_outputs/final_revised_draft/Messinger_Dissertation_Final_Revised_Draft.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


doc = Document(SOURCE)

# Remove template sample reference/appendix/caption instructions that follow the real reference list.
cut_at = None
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text in {"APA Reference Sample Entries", "Sample Entries"}:
        cut_at = i
        break

if cut_at is not None:
    for para in list(doc.paragraphs[cut_at:]):
        delete_paragraph(para)

while doc.paragraphs and not doc.paragraphs[-1].text.strip():
    delete_paragraph(doc.paragraphs[-1])

# Enforce double spacing throughout the revised dissertation draft.
for style_name in [
    "Normal",
    "Body Text",
    "Normal (Web)",
    "List Paragraph",
    "APA Level 0",
    "APA Level 1",
    "APA Level 2",
    "APA Level 3",
    "APA References",
    "APA Reference",
]:
    if style_name in doc.styles:
        style = doc.styles[style_name]
        style.paragraph_format.line_spacing = 2.0
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)

for para in doc.paragraphs:
    para.paragraph_format.line_spacing = 2.0
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.line_spacing = 2.0
                para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(0)

doc.save(OUT)
print(OUT)
