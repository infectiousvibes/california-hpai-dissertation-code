from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


SOURCE = Path("analysis_outputs/chapter5_draft/Messinger_Dissertation_Chapters4_5_Draft.docx")
OUT = Path("analysis_outputs/dissertation_revisions/Messinger_Dissertation_Revised_WildBird_NullFindings.docx")


def set_text(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


def insert_after(paragraph, text, style=None):
    new_p = deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.clear()
    if style:
        inserted.style = style
    inserted.add_run(text)
    return inserted


doc = Document(SOURCE)

# Abstract / front matter
set_text(
    doc.paragraphs[19],
    "Highly Pathogenic Avian Influenza (HPAI) H5N1 continues to threaten agriculture, wildlife, and public health across California. Flooded rice fields in the Central Valley increasingly function as surrogate wetlands for migratory waterfowl, particularly Lesser Snow Geese, creating potential ecological interfaces where wild-bird movement, managed agricultural habitat, and domestic poultry production may converge. This quantitative ecological study examined whether county-level rice acreage, Lesser Snow Goose activity, and their interaction were associated with poultry HPAI H5N1 outbreak incidence during the complete winter seasons available for analysis. Guided by One Health and disease ecology frameworks, the study treated wild-bird migration and rice-field habitat as the primary ecological exposure pathway and poultry outbreak incidence as the measurable spillover outcome. Secondary data were drawn from USDA, eBird, PRISM, and public surveillance sources and analyzed using negative binomial regression, logistic sensitivity models, and GIS-based spatial overlays. The adjusted regression models did not identify rice acreage, Snow Goose activity, or their interaction as statistically significant independent predictors of poultry outbreak counts. However, GIS outputs showed spatial convergence among rice agriculture, Snow Goose activity, poultry inventory, wild-bird detections, and outbreak burden in selected counties. These findings support a cautious interpretation: ecological convergence remains theoretically plausible, but county-level archival data may be too coarse to detect fine-scale wildlife-poultry transmission pathways. The study contributes to One Health surveillance by clarifying both the value and limits of county-level ecological modeling for HPAI preparedness.",
)

# Chapter 1: purpose, RQs, conceptual model, nature of study
set_text(
    doc.paragraphs[146],
    "This quantitative ecological study uses county-level data to examine relationships among post-harvest rice acreage, Lesser Snow Goose (Anser caerulescens) activity, and poultry HPAI H5N1 outbreak incidence in California’s Central Valley and related Pacific Flyway agricultural-wetland landscapes from 2020 to 2025. The study is centered on the wild-bird migration and agricultural-habitat pathway; poultry outbreak incidence is used as the measurable endpoint of potential spillover into domestic animal systems. The unit of analysis is the California county-winter season that meets the study’s inclusion criteria. By quantifying associations between rice-field habitat, wild-bird activity, and outbreak locations, both individually and jointly, the study seeks to identify ecological conditions that may be linked to viral amplification and spillover at the agroecological interface. In doing so, it addresses a critical knowledge gap at the intersection of agricultural land use, migratory bird ecology, and emerging infectious diseases and supports more targeted, integrated One Health strategies.",
)

rq_replacements = {
    148: "Research Question 1 (RQ1) What is the relationship between rice acreage as a proxy for potential flooded rice habitat and poultry HPAI H5N1 outbreak incidence in California counties within the Central Valley and related Pacific Flyway agricultural-wetland interface between 2020 and 2025?",
    150: "Null Hypothesis (H01): There is no statistically significant relationship between rice acreage and poultry HPAI H5N1 outbreak incidence in California counties within the Central Valley and related Pacific Flyway agricultural-wetland interface between 2020 and 2025.",
    152: "Alternative Hypothesis (H1₁): There is a statistically significant relationship between rice acreage and poultry HPAI H5N1 outbreak incidence in California counties within the Central Valley and related Pacific Flyway agricultural-wetland interface between 2020 and 2025.",
    154: "Research Question 2 (RQ2) What is the relationship between Lesser Snow Goose activity as an indicator of wild-bird migration and poultry HPAI H5N1 outbreak incidence in California counties within the Central Valley and related Pacific Flyway agricultural-wetland interface from 2020 to 2025?",
    156: "Null Hypothesis (H02): Lesser Snow Goose activity is not significantly associated with poultry HPAI H5N1 outbreak incidence in California counties within the Central Valley and related Pacific Flyway agricultural-wetland interface from 2020 to 2025.",
    157: "Alternative Hypothesis (H1₂): Lesser Snow Goose activity is significantly associated with poultry HPAI H5N1 outbreak incidence in California counties within the Central Valley and related Pacific Flyway agricultural-wetland interface from 2020 to 2025.",
    159: "Research Question 3 (RQ3) Does the relationship between rice acreage and poultry HPAI H5N1 outbreak incidence vary by Lesser Snow Goose activity in California counties within the Central Valley and related Pacific Flyway agricultural-wetland interface between 2020 and 2025?",
    161: "Null Hypothesis (H03): The relationship between rice acreage and poultry HPAI H5N1 outbreak incidence does not vary by Lesser Snow Goose activity in California counties within the Central Valley and related Pacific Flyway agricultural-wetland interface between 2020 and 2025.",
    163: "Alternative Hypothesis (H1₃): The relationship between rice acreage and poultry HPAI H5N1 outbreak incidence varies by Lesser Snow Goose activity in California counties within the Central Valley and related Pacific Flyway agricultural-wetland interface between 2020 and 2025.",
}
for idx, text in rq_replacements.items():
    set_text(doc.paragraphs[idx], text)

set_text(
    doc.paragraphs[167],
    "Using these theories, the study views HPAI H5N1 not only as a poultry-industry issue but also as a complex environmental and public health challenge driven by land use, ecological pressures, wild-bird movement, and agricultural exposure pathways (Patz et al., 2004; Wu et al., 2020; Zinsstag et al., 2011). This theoretical approach guides the selection of landscape-level predictors, such as rice-field habitat and waterfowl activity, while treating poultry inventory as the susceptible domestic host context through which spillover is observed (Harvey et al., 2023; McDuie et al., 2022; Overton & Casazza, 2023).",
)
set_text(doc.paragraphs[169], "Building on this literature, the framework suggests that four interconnected forces shape observed poultry HPAI H5N1 outbreak burden in California:")
set_text(
    doc.paragraphs[174],
    "Taken together, these components form a conceptual framework in which anthropogenic habitat modification and climate variability shape wild-bird aggregation, while poultry exposure influences where ecological exposure may become an observed domestic-animal outbreak. By modeling these factors together, this study aims to quantify the extent to which agricultural habitat and wild-bird activity predict outbreak risk across space and time after accounting for susceptible poultry populations (Finucane et al., 2014; Harvey et al., 2023).",
)
set_text(
    doc.paragraphs[177],
    "This study uses a quantitative, retrospective ecological research design with county-level archival data from 2020 to 2025 in California’s Central Valley and related Pacific Flyway agricultural-wetland landscapes. The unit of analysis is the county-winter season. The study draws on secondary data sources to examine whether rice acreage as a proxy for potential flooded rice habitat and Lesser Snow Goose activity as a wild-bird migration indicator are statistically associated with poultry HPAI H5N1 outbreak incidence.",
)
set_text(
    doc.paragraphs[188],
    "Data will be analyzed using multivariable generalized linear models appropriate for rate outcomes, including interaction terms to evaluate whether the association between rice acreage and poultry outbreak incidence varies by Lesser Snow Goose activity. Poultry inventory will be included as a susceptible host covariate rather than as the main ecological exposure. GIS-based spatial overlays will be used to describe and visualize spatial convergence between rice-field habitat, waterfowl distributions, poultry exposure, wild-bird detections, and HPAI H5N1 outbreak incidence, thereby supporting interpretation of results for RQ1-RQ3.",
)

# Chapter 3: design, population/sample, variables, and summary
set_text(
    doc.paragraphs[363],
    "This study uses a retrospective, quantitative ecological design to examine relationships between agricultural land use, wild-bird migration indicators, and the spatial distribution of poultry HPAI H5N1 outbreaks in California from 2020 to 2025. The primary goal is to determine whether two ecological variables, rice acreage as a proxy for potential post-harvest flooded rice habitat and Lesser Snow Goose (Anser caerulescens) activity, are significantly associated with county-level poultry HPAI H5N1 outbreak incidence and whether their interaction is related to elevated outbreak risk. Poultry inventory is included as a covariate representing the susceptible domestic host population. This approach is suitable because the unit of analysis is ecological (county-winter season) and the data are obtained from existing secondary sources (Chanda et al., 2025; Harvey et al., 2023).",
)
set_text(
    doc.paragraphs[370],
    "The unit of analysis is the California county-winter season, and the methodology adopts a spatial-temporal framework consistent with landscape epidemiology and disease ecology approaches (Patz et al., 2004; Zinsstag et al., 2011). All independent and dependent variables are measured at the county-winter-season level (November-February), aligning with the migratory and agricultural cycles under investigation. The key ecological predictors, rice acreage and Lesser Snow Goose activity, are analyzed in relation to poultry HPAI H5N1 outbreak incidence in each county-winter season (Martinico et al., 2024; Overton & Casazza, 2023; Teitelbaum et al., 2024). Data were extracted, cleaned, coded, and merged using R and GIS tools, including QGIS-compatible spatial outputs, which are widely used for spatial analyses of infectious disease risk (Chanda et al., 2025; Wu et al., 2020).",
)
set_text(
    doc.paragraphs[374],
    "The population for this study consists of California county-winter observations that meet ecological and agricultural criteria relevant to the research questions. In an ecological study, the “population” refers to spatial units rather than individual people, birds, or farms; in this case, each county-winter season constitutes a single observational unit (Patz et al., 2004; Zinsstag et al., 2011). This place-based focus aligns with landscape epidemiology and disease ecology frameworks that emphasize geographic determinants of zoonotic risk.",
)
set_text(doc.paragraphs[375], "The target population includes California county-winter observations within the Pacific Flyway agricultural-wetland interface that:")
set_text(doc.paragraphs[376], "Intersect the Central Valley or related agricultural-wetland landscapes used by migratory waterfowl;")
set_text(doc.paragraphs[377], "Have rice agriculture, managed wetland habitat, or documented migratory waterfowl activity relevant to the wild-bird exposure pathway; and")
set_text(doc.paragraphs[378], "Have available data for poultry HPAI H5N1 outbreak outcomes, wild-bird activity or surveillance, poultry exposure, and climate covariates for at least one winter season during the study period.")
set_text(
    doc.paragraphs[379],
    "This population was selected because California’s Central Valley is both a major rice-growing region and an important wintering habitat for migratory waterfowl, including the Lesser Snow Goose (Baldwin et al., 2022; Martinico et al., 2024; Takekawa et al., 2023). These counties are where wildlife, agricultural land, and domestic poultry production converge, conditions that have been linked to pathogen spillover and avian influenza ecology in prior work (Harvey et al., 2023; Teitelbaum et al., 2024). For the primary inferential analysis, the analytic sample was narrowed to rice-producing counties with at least 1,000 USDA Cropland Data Layer rice acres in any crop year from 2022 through 2024 and complete winter-season covariate data. A statewide county-winter sensitivity analysis was retained to evaluate whether findings were robust beyond the rice-producing county sample.",
)
set_text(
    doc.paragraphs[384],
    "Although this study does not involve human participants or experimental sampling, it uses a defined sampling frame based on counties as ecological units. A clearly articulated sampling strategy is therefore necessary to specify how county-winter observations enter the analysis (Patz et al., 2004; Zinsstag et al., 2011). The approach used here is best described as purposive and census-based: eligible county-winter observations are included according to predefined ecological, agricultural, and data-completeness criteria rather than random sampling.",
)
set_text(doc.paragraphs[385], "County-winter observations were included in the primary analysis if they satisfied the following criteria:")
set_text(doc.paragraphs[386], "Were located within California’s Central Valley or related Pacific Flyway agricultural-wetland interface;")
set_text(doc.paragraphs[387], "Had at least 1,000 USDA Cropland Data Layer rice acres in any crop year from 2022 through 2024, with DWR Statewide Crop Mapping used as a supplemental validation source where available (Matthews et al., 2022; Rosenberg et al., 2022; Strum et al., 2013);")
set_text(doc.paragraphs[388], "Had available poultry HPAI H5N1 outbreak outcome data from USDA APHIS or California animal-health surveillance sources (USDA APHIS, 2023; Kan-Rice, 2025);")
set_text(doc.paragraphs[389], "Had observable migratory waterfowl activity, especially Lesser Snow Goose activity, documented through eBird summaries, with APHIS wild-bird detections retained as a supplemental surveillance indicator (Baldwin et al., 2022; McDuie et al., 2022, 2024; Sullivan et al., 2025).")
set_text(
    doc.paragraphs[390],
    "No eligible counties were randomly excluded. Counties that did not meet the rice-acreage threshold were excluded from the primary rice-producing county analysis but retained in the statewide sensitivity dataset when complete covariate and outcome data were available. This distinction preserves the Pacific Flyway and wild-bird migration rationale while making the final analytic sample transparent and reproducible.",
)
set_text(
    doc.paragraphs[424],
    "Poultry HPAI H5N1 outbreak incidence.\nThe primary outcome is poultry outbreak incidence per county-winter season, defined as the number of confirmed HPAI H5N1 outbreaks in domestic poultry within each county-winter season. Outbreaks are identified using laboratory-confirmed detections reported by USDA APHIS and CDFA (Green et al., 2023; USDA APHIS, 2023). For secondary analyses, a binary measure (0 = no poultry outbreak, 1 = ≥1 poultry outbreak in the county-winter season) is used to compare the presence versus absence of poultry outbreaks over time and space (Chanda et al., 2025; Mena et al., 2025).",
)
set_text(
    doc.paragraphs[427],
    "Lesser Snow Goose activity.\nThis variable represents the relative winter activity of Lesser Snow Geese (Anser caerulescens) in each county-winter season and serves as the primary wild-bird migration indicator for RQ2 and RQ3. Goose observations are obtained from eBird.org and filtered by species, date (November-February), and county. To account for differences in observer effort across counties, activity is standardized as the mean number of Lesser Snow Geese reported per complete eBird checklist per county-winter season (Baldwin et al., 2022; Carter et al., 2025). This measure provides a relative index of species presence and aggregation while partially controlling for variation in reporting effort. Spatial and temporal patterns are interpreted alongside telemetry and movement literature documenting major staging and wintering areas (McDuie et al., 2022, 2024; Overton & Casazza, 2023; Sullivan et al., 2025).",
)
set_text(
    doc.paragraphs[433],
    "Poultry operation and inventory exposure.\nDerived from the USDA Agricultural Census, this variable reflects susceptible domestic host exposure in each county. Poultry inventory and operation measures are included as covariates because counties with more poultry have greater opportunity for observed poultry outbreaks. They are not the main ecological exposure of interest; rather, they help distinguish wild-bird and rice-field habitat associations from the underlying distribution of domestic poultry populations (Green et al., 2023; Patyk et al., 2023).",
)
set_text(
    doc.paragraphs[447],
    "RQ1: How does variation in rice acreage as a proxy for potential flooded rice habitat relate to poultry HPAI H5N1 outbreak incidence in California counties within the Central Valley and related Pacific Flyway agricultural-wetland interface between 2020 and 2025?\nH₀₁: There is no statistically significant relationship between rice acreage and poultry HPAI H5N1 outbreak incidence.\nH₁₁: There is a statistically significant relationship between rice acreage and poultry HPAI H5N1 outbreak incidence.",
)
set_text(
    doc.paragraphs[448],
    "RQ2: How does variation in Lesser Snow Goose activity, as an indicator of wild-bird migration and aggregation, correspond with poultry HPAI H5N1 outbreak incidence in California counties within the Central Valley and related Pacific Flyway agricultural-wetland interface from 2020 to 2025?\nH₀₂: There is no statistically significant relationship between Lesser Snow Goose activity and poultry HPAI H5N1 outbreak incidence.\nH₁₂: There is a statistically significant relationship between Lesser Snow Goose activity and poultry HPAI H5N1 outbreak incidence.",
)
set_text(
    doc.paragraphs[449],
    "RQ3: Does the relationship between rice acreage and poultry HPAI H5N1 outbreak incidence vary depending on Lesser Snow Goose activity in California counties between 2020 and 2025?\nH₀₃: The relationship between rice acreage and poultry HPAI H5N1 outbreak incidence does not vary depending on Lesser Snow Goose activity.\nH₁₃: The relationship between rice acreage and poultry HPAI H5N1 outbreak incidence does vary depending on Lesser Snow Goose activity.",
)
set_text(
    doc.paragraphs[496],
    "This chapter details the research methodology used to examine the relationship between environmental and ecological variables, specifically rice acreage and Lesser Snow Goose activity, and the occurrence of poultry HPAI H5N1 outbreaks across California county-winter observations between 2020 and 2025. A non-experimental, ecological study design was chosen to analyze population-level associations using archival data from publicly available sources.",
)
set_text(
    doc.paragraphs[497],
    "The methodology specifies California counties as the unit of analysis, outlines inclusion criteria, and describes the use of a purposive census approach rather than random sampling. Key variables are operationalized using validated sources, including the USDA Cropland Data Layer, eBird, APHIS wild-bird and poultry outbreak surveillance, PRISM climate data, and USDA Agricultural Census poultry covariates. Negative binomial regression is used as the primary count-modeling approach, with logistic regression retained as a secondary binary-outcome analysis. Spatial overlays are used to interpret convergence among rice-field habitat, wild-bird activity, poultry exposure, and outbreak burden.",
)

# Chapter 4: align results interpretation with revised framing.
set_text(
    doc.paragraphs[502],
    "The purpose of this quantitative, retrospective ecological study was to examine whether county-level rice acreage and Lesser Snow Goose activity, representing agricultural habitat and wild-bird migration pathways, were associated with poultry HPAI H5N1 outbreak incidence in California during the complete winter seasons available for analysis. The chapter presents the data collection results, descriptive findings, statistical assumption checks, model results organized by research question, supplemental wild-bird surveillance indicators, and GIS-based spatial convergence findings.",
)
set_text(
    doc.paragraphs[503],
    "The three research questions were: (a) whether rice acreage was associated with poultry HPAI H5N1 outbreak incidence, (b) whether Lesser Snow Goose activity was associated with poultry HPAI H5N1 outbreak incidence, and (c) whether the relationship between rice acreage and poultry outbreak incidence varied by Lesser Snow Goose activity. The corresponding null hypotheses stated that there would be no statistically significant association for rice acreage, Snow Goose activity, or their interaction. Poultry inventory was included as an adjustment covariate because it represents susceptible domestic host exposure, not because the study was conceptually centered only on poultry production.",
)
set_text(
    doc.paragraphs[547],
    "Research Question 2 asked whether Lesser Snow Goose activity, the primary wild-bird migration indicator, was associated with poultry HPAI H5N1 outbreak incidence. In the primary negative binomial model, the Snow Goose metric was not statistically significant after adjustment for poultry inventory, eBird checklist effort, winter temperature, and winter precipitation, RR = 1.10, 95% CI [0.60, 2.03], p = .762. The null hypothesis for RQ2 was not rejected. The statewide sensitivity model also did not show a statistically significant adjusted association between the Snow Goose metric and outbreak count, RR = 1.02, 95% CI [0.64, 1.62], p = .943. Therefore, the wild-bird migration hypothesis remained theoretically plausible but was not statistically supported by the adjusted county-winter models.",
)
set_text(
    doc.paragraphs[560],
    "The GIS analysis was used to evaluate spatial convergence among rice acreage, Snow Goose activity, poultry inventory, wild-bird surveillance indicators, and outbreak burden. This analysis was complementary to the regression models: the regression models evaluated adjusted statistical associations, whereas the maps evaluated where ecological and agricultural risk components overlapped geographically. The One Health spatial convergence index combined ranked county values for rice acreage, Snow Goose activity, and poultry inventory, with wild-bird detections interpreted as supplemental surveillance context.",
)
set_text(
    doc.paragraphs[596],
    "This chapter presented descriptive, statistical, supplemental wild-bird surveillance, and GIS findings for the county-winter ecological analysis. Outbreak counts increased across complete winter seasons, with the largest burden observed in 2024-2025. Negative binomial regression was selected for the primary count models because the outcome was sparse and overdispersed. The adjusted models did not identify CDL rice acreage, Snow Goose activity, or their interaction as statistically significant independent predictors of poultry HPAI outbreak counts. Therefore, the null hypotheses for RQ1, RQ2, and RQ3 were not rejected.",
)
set_text(
    doc.paragraphs[597],
    "Although the focal ecological predictors were not statistically significant in the adjusted regression models, GIS analysis showed that rice acreage, Snow Goose activity, poultry inventory, wild-bird detections, and outbreak burden exhibited spatial convergence in selected Central Valley counties. The findings suggest that county-level rice acreage and Snow Goose metrics may indicate potential ecological risk zones but may be too coarse to capture the finer-scale wildlife-poultry interfaces where transmission risk is most likely to occur. Chapter 5 interprets these findings in relation to the One Health and disease ecology frameworks, discusses limitations of county-level archival data, and presents recommendations for future research and surveillance.",
)

# Chapter 5: sharpen the null-finding interpretation around wild-bird exposure and poultry outcome.
set_text(
    doc.paragraphs[611],
    "Research Question 2 asked whether Lesser Snow Goose activity, the primary wild-bird migration indicator, was associated with poultry HPAI H5N1 outbreak incidence. The adjusted negative binomial models did not identify the eBird Snow Goose metric as a statistically significant independent predictor in either the primary rice-producing county dataset or the statewide sensitivity dataset. Therefore, the null hypothesis for RQ2 was not rejected.",
)
set_text(
    doc.paragraphs[613],
    "The Snow Goose finding also suggests that species-specific abundance may not capture the broader wild-bird exposure environment. HPAI dynamics involve multiple wild bird species, environmental persistence, mortality surveillance, and changing viral ecology (Criado et al., 2025; Harvey et al., 2023). Supplemental APHIS wild-bird detections were therefore useful as descriptive surveillance context, but they were not used to replace the eBird Snow Goose metric as the focal exposure because detection records are strongly influenced by surveillance effort and reporting. Snow Geese remain theoretically important because of their large wintering flocks and documented relevance to California agricultural landscapes, but poultry outbreak risk may reflect a more complex waterfowl and wild-bird community rather than a single species indicator.",
)
set_text(
    doc.paragraphs[619],
    "The findings align with the One Health framework by showing that poultry HPAI outbreak patterns cannot be understood through a single ecological or agricultural variable alone. One Health emphasizes interconnected human, animal, and environmental systems (Zinsstag et al., 2011). In this study, wild-bird migration and rice-field habitat formed the central ecological exposure pathway, while poultry inventory represented the susceptible domestic host context through which spillover was observed. Poultry inventory emerged as the most consistent statistical predictor, while GIS maps showed that ecological and agricultural risk components converged spatially in selected counties. Together, these findings suggest that HPAI risk is best understood as a coupled system involving domestic poultry exposure, wild-bird ecology, agricultural landscapes, and climate conditions.",
)
set_text(
    doc.paragraphs[622],
    "The results extend prior literature by applying an integrated county-winter analysis to California HPAI H5N1 outbreak patterns. Prior work has demonstrated that wild waterfowl use agricultural landscapes and may create pathways for viral exposure near livestock and poultry systems (McDuie et al., 2022a; Teitelbaum et al., 2024). Other studies have emphasized the importance of wetlands, agricultural habitat, and land-use patterns in shaping avian influenza risk (Overton & Casazza, 2023; Wu et al., 2020). The current study adds that these relationships may not be detectable as simple independent effects when measured through county-level annual rice acreage, county-winter bird observation summaries, and surveillance-dependent wild-bird detection records.",
)
set_text(
    doc.paragraphs[644],
    "Methodologically, the study demonstrates both the usefulness and limits of publicly available archival data for ecological disease research. The integrated dataset allowed county-winter modeling across agricultural habitat, wild-bird activity, poultry exposure, climate, and outbreak outcomes, but the non-significant focal predictors show that publicly available county-level proxies may be insufficient for detecting fine-scale wildlife-poultry transmission pathways. This finding supports the need for more precise spatial data and better integration of agricultural, wildlife, and environmental datasets.",
)
set_text(
    doc.paragraphs[651],
    "This study examined whether rice acreage, Lesser Snow Goose activity, and their interaction were associated with poultry HPAI H5N1 outbreak incidence in California using county-winter archival data. The adjusted regression models did not find statistically significant independent associations for rice acreage, Snow Goose activity, or their interaction. The null hypotheses for all three research questions were therefore not rejected. At the same time, outbreak burden increased substantially across the study period, poultry inventory emerged as the most consistent predictor, supplemental wild-bird detections supported the relevance of wildlife surveillance context, and GIS analysis identified counties where ecological and agricultural risk components converged spatially.",
)
set_text(
    doc.paragraphs[652],
    "The central conclusion is that the theory of ecological convergence remains plausible, but the county-level archival data used in this study were not precise enough to demonstrate the hypothesized independent effects statistically. Rice acreage and Snow Goose activity may indicate potential ecological exposure, while wild-bird detections provide supplemental evidence of detected wildlife viral activity; however, observed poultry outbreak occurrence appears to depend on poultry density and finer-scale contact pathways that were not fully captured in the available data. The study contributes to the field by showing both the value and the limits of county-level One Health modeling for HPAI risk, and by pointing toward the need for finer-scale, integrated surveillance of rice-field habitat, wild-bird movement, poultry exposure, and environmental conditions.",
)

# Insert new paragraphs after all fixed-index replacements. Insert from the bottom upward
# so earlier insertions do not shift later targets.
insert_after(
    doc.paragraphs[523],
    "Wild-bird surveillance variables were also summarized to retain the migration and wildlife component of the study design. In the primary rice-producing county dataset, APHIS wild-bird detections averaged 7.28 detections per county-winter observation (SD = 11.75), and in the statewide sensitivity dataset they averaged 2.89 detections (SD = 6.50). Supplemental rank screening showed weak positive relationships between wild-bird detections and poultry outbreak counts in both the primary dataset (Spearman ρ = .253) and the statewide sensitivity dataset (Spearman ρ = .244). These descriptive results support the relevance of the wild-bird surveillance pathway, but they were interpreted cautiously because APHIS wild-bird detections are influenced by surveillance intensity and reporting practices.",
    style=doc.paragraphs[523].style,
)
insert_after(
    doc.paragraphs[427],
    "Supplemental wild-bird detections.\nAPHIS wild-bird and waterfowl detections are retained as descriptive and sensitivity indicators of detected local viral activity. These detections are not treated as the primary RQ2 exposure because they are influenced by surveillance intensity, carcass discovery, reporting practices, and outbreak-response sampling. They are useful for triangulating the wild-bird migration pathway but are interpreted more cautiously than the eBird-based Snow Goose activity measure.",
    style=doc.paragraphs[427].style,
)
insert_after(
    doc.paragraphs[172],
    "Poultry inventory and operation context as indicators of susceptible domestic host exposure (Green et al., 2023; Patyk et al., 2023)",
    style=doc.paragraphs[172].style,
)

doc.save(OUT)
print(OUT)
