from pathlib import Path

from docx import Document
from docx.shared import Pt


CHECKLIST = Path("analysis_outputs/final_revised_draft/Messinger_Quantitative_Checklist_Final_Revised_Draft.docx")
OUT = CHECKLIST


def set_cell_paragraphs(cell, paragraphs):
    cell.text = ""
    # python-docx leaves one paragraph after clearing text.
    first = True
    for text in paragraphs:
        para = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        para.text = text
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(3)
        para.paragraph_format.line_spacing = 1.0
        for run in para.runs:
            run.font.size = Pt(9)


def comment_cell(table, row_idx):
    return table.rows[row_idx].cells[-1]


def set_comment(table, row_idx, paragraphs):
    set_cell_paragraphs(comment_cell(table, row_idx), paragraphs)


def set_pages(table, rows, page):
    for row_idx in rows:
        cells = table.rows[row_idx].cells
        if len(cells) >= 3:
            cells[-2].text = page


doc = Document(CHECKLIST)

# Front matter.
set_comment(doc.tables[0], 3, [
    "Type of relationship: The study is correlational/ecological, examining county-winter associations among agricultural habitat, wild-bird activity, poultry exposure, climate, and poultry HPAI outbreak outcomes.",
    "Key variables: Independent variables are rice acreage as a proxy for potential flooded rice habitat and Lesser Snow Goose activity as the wild-bird migration indicator. The dependent variable is poultry HPAI H5N1 outbreak incidence. Covariates include poultry inventory, eBird checklist effort, winter temperature, and winter precipitation.",
    "Applicable population: The target population is California county-winter observations in the Pacific Flyway agricultural-wetland interface, with a primary rice-producing county sample and statewide sensitivity dataset.",
])
for row in [4, 5]:
    set_comment(doc.tables[0], row, doc.tables[0].rows[3].cells[-1].text.split("\n"))

set_comment(doc.tables[0], 7, [
    "Research problem: The abstract identifies HPAI H5N1 as a wildlife, agricultural, and public health concern in California’s Central Valley and related Pacific Flyway landscapes.",
    "Purpose: The abstract states that the study examined whether rice acreage, Lesser Snow Goose activity, and their interaction were associated with poultry HPAI outbreak incidence.",
    "Framework: The abstract identifies One Health and disease ecology as the guiding frameworks for interpreting ecological spillover risk.",
    "Research questions: The abstract summarizes RQ1-RQ3 around rice acreage, Snow Goose activity, and the rice by Snow Goose interaction.",
    "Design/methods: The abstract describes a quantitative retrospective ecological design using archival USDA, eBird, PRISM, and surveillance data with regression and GIS analyses.",
    "Key findings: The abstract reports that the adjusted models did not find statistically significant focal predictors, while GIS outputs showed spatial convergence in selected counties.",
    "Positive social change: The abstract connects findings to targeted surveillance, preparedness, agricultural planning, and One Health coordination.",
])
for row in range(8, 14):
    set_comment(doc.tables[0], row, doc.tables[0].rows[7].cells[-1].text.split("\n"))

# Chapter 1.
set_comment(doc.tables[1], 4, [
    "Topic/need/social change: The introduction describes HPAI H5N1 as a growing One Health concern involving wild birds, rice-field habitat, poultry exposure, agriculture, and public health preparedness.",
    "Chapter preview: The introduction previews the problem, purpose, research questions, theoretical/conceptual framework, nature of the study, definitions, assumptions, scope, limitations, and significance.",
])
set_comment(doc.tables[1], 5, doc.tables[1].rows[4].cells[-1].text.split("\n"))

set_comment(doc.tables[2], 1, [
    "Literature scope: The background summarizes current literature on HPAI H5N1, migratory waterfowl, Central Valley rice agriculture as surrogate wetland habitat, climate variability, and wildlife-agriculture interfaces.",
    "Gap in knowledge: The section identifies limited empirical research quantifying how rice-field habitat and wild-bird activity relate to HPAI outbreak patterns in California.",
    "Why needed: The background ends by explaining that integrated ecological analysis is needed to support proactive surveillance and One Health risk reduction.",
])
for row in [2, 3]:
    set_comment(doc.tables[2], row, doc.tables[2].rows[1].cells[-1].text.split("\n"))

set_comment(doc.tables[2], 5, [
    "Research problem: The problem is the limited understanding of how rice-field habitat, migratory waterfowl activity, and poultry exposure are associated with HPAI outbreak patterns.",
    "Current/relevant/significant: The section uses recent HPAI spread, wildlife detections, poultry outbreaks, and One Health concerns to show that the issue is current and significant.",
    "Connection to recent research: The problem builds from recent studies on HPAI ecology, waterfowl movement, surrogate wetlands, and poultry exposure pathways.",
    "Meaningful gap: The gap is the absence of integrated, California-specific, county-level analyses linking ecological predictors and outbreak outcomes.",
])
for row in [6, 7, 8]:
    set_comment(doc.tables[2], row, doc.tables[2].rows[5].cells[-1].text.split("\n"))

set_comment(doc.tables[2], 10, [
    "Quantitative study: The purpose statement explicitly identifies the study as quantitative and ecological.",
    "Intent/variables: The study intent is to examine associations among rice acreage, Lesser Snow Goose activity, their interaction, and poultry HPAI outbreak incidence, while accounting for poultry inventory, eBird effort, and climate covariates.",
])
set_comment(doc.tables[2], 11, doc.tables[2].rows[10].cells[-1].text.split("\n"))

set_comment(doc.tables[2], 13, [
    "Research questions: RQ1 asks about rice acreage and poultry outbreak incidence; RQ2 asks about Lesser Snow Goose activity and poultry outbreak incidence; RQ3 asks whether the rice-outbreak relationship varies by Snow Goose activity.",
    "Hypotheses: The null and alternative hypotheses identify the independent variables, dependent variable, and tested associations/interactions for each research question.",
])
set_comment(doc.tables[2], 14, doc.tables[2].rows[13].cells[-1].text.split("\n"))

set_comment(doc.tables[3], 2, [
    "Theory/source: The theoretical foundation identifies One Health, which frames human, animal, and environmental health as interconnected.",
    "Major propositions: One Health supports the proposition that HPAI risk arises through linked environmental habitat, wild-bird movement, domestic animal exposure, and public health systems.",
    "Relation to study: The theory supports examining rice-field habitat, wild-bird activity, and poultry exposure together rather than treating outbreak risk as poultry-only.",
])
for row in [3, 4]:
    set_comment(doc.tables[3], row, doc.tables[3].rows[2].cells[-1].text.split("\n"))

set_comment(doc.tables[3], 6, [
    "Concept/phenomenon: The conceptual framework is an eco-epidemiological model of HPAI spillover at the wildlife-agriculture interface.",
    "Literature basis: The framework is derived from literature on wetlands, rice agriculture, migratory waterfowl, environmental persistence, poultry exposure, and One Health surveillance.",
    "Logical connections: Rice acreage indicates potential habitat, Snow Goose activity indicates wild-bird aggregation, poultry inventory indicates susceptible host exposure, and climate modifies habitat and virus persistence.",
    "Relation to methods: The framework directly informs the variables, interaction term, covariates, regression models, and GIS convergence outputs.",
])
for row in [7, 8, 9]:
    set_comment(doc.tables[3], row, doc.tables[3].rows[6].cells[-1].text.split("\n"))

set_comment(doc.tables[3], 11, [
    "Design rationale: A retrospective ecological design is appropriate because the research questions focus on county-winter population-level associations rather than individual exposures.",
    "Key variables: IVs are rice acreage and Lesser Snow Goose activity; the DV is poultry HPAI outbreak incidence; covariates are poultry inventory, eBird checklist effort, temperature, and precipitation.",
    "Method summary: Data are drawn from public archival sources and analyzed using regression models and GIS spatial overlays.",
])
for row in [12, 13]:
    set_comment(doc.tables[3], row, doc.tables[3].rows[11].cells[-1].text.split("\n"))

set_comment(doc.tables[4], 1, [
    "Variable definitions: The section defines rice acreage/flooded rice habitat, Lesser Snow Goose activity, poultry HPAI outbreak incidence, poultry exposure, and climate covariates.",
    "Multiple-meaning terms: Definitions clarify how ecological, agricultural, and surveillance terms are used in this study so they are not interpreted as individual-level exposure claims.",
    "Literature support: Definitions are supported by literature and data-source descriptions, with more detailed coding and operationalization in Chapter 3.",
])
for row in [2, 3]:
    set_comment(doc.tables[4], row, doc.tables[4].rows[1].cells[-1].text.split("\n"))

set_comment(doc.tables[4], 5, [
    "Critical assumptions: The study assumes that archival USDA, eBird, PRISM, APHIS/CDFA, and Agricultural Census data are sufficiently accurate for county-level ecological analysis.",
    "Why necessary: These assumptions are necessary because the study uses secondary data and proxy measures rather than direct field observation of waterfowl, flooding, farms, or virus movement.",
])
set_comment(doc.tables[4], 6, doc.tables[4].rows[5].cells[-1].text.split("\n"))

set_comment(doc.tables[4], 8, [
    "Specific focus: The study focuses on ecological correlates of HPAI outbreak incidence, especially rice-field habitat and wild-bird activity in California agricultural-wetland landscapes.",
    "Boundaries: Included populations are California county-winter observations with relevant ecological/agricultural data; excluded are individual birds, farms, people, and non-California systems.",
    "Generalizability: Findings may apply most directly to similar agricultural-wetland and migratory waterfowl interface regions, not all poultry or wildlife systems.",
])
for row in [9, 10]:
    set_comment(doc.tables[4], row, doc.tables[4].rows[8].cells[-1].text.split("\n"))

set_comment(doc.tables[5], 1, [
    "Design/method limitations: Limitations include the ecological county-level design, small primary sample, sparse outcomes, and inability to infer farm-level or individual-level causation.",
    "Biases: Potential biases include eBird observer effort, surveillance/reporting differences in wild-bird detections, static poultry inventory data, and measurement error in rice acreage as a flooding proxy.",
    "Measures to address: The study uses covariate adjustment, effort measures, negative binomial models, statewide sensitivity analysis, and cautious interpretation of null findings.",
])
for row in [2, 3]:
    set_comment(doc.tables[5], row, doc.tables[5].rows[1].cells[-1].text.split("\n"))

set_comment(doc.tables[5], 5, [
    "Advance knowledge: The study contributes integrated ecological evidence about HPAI risk at the rice-field, wild-bird, and poultry interface.",
    "Advance practice/policy: Findings may support targeted One Health surveillance and coordination among wildlife, agriculture, and public health partners.",
    "Positive social change: Social change implications are bounded to improved preparedness, food-system protection, and prevention-focused planning.",
])
for row in [6, 7]:
    set_comment(doc.tables[5], row, doc.tables[5].rows[5].cells[-1].text.split("\n"))

set_comment(doc.tables[5], 9, [
    "Chapter summary: Chapter 1 summarizes the topic, problem, purpose, RQs, hypotheses, frameworks, scope, assumptions, limitations, and significance.",
    "Transition: The chapter transitions to Chapter 2 by identifying the need to review literature on HPAI ecology, waterfowl movement, rice agriculture, poultry exposure, and One Health.",
])
set_comment(doc.tables[5], 10, doc.tables[5].rows[9].cells[-1].text.split("\n"))

# Chapter 2.
set_comment(doc.tables[6], 3, [
    "Problem/purpose restated: Chapter 2 restates the need to understand ecological and agricultural factors associated with HPAI H5N1 outbreak patterns.",
    "Literature relevance: The introduction explains that current HPAI expansion, wild-bird movement, and agricultural habitat changes make the problem relevant.",
    "Preview: The chapter previews search strategy, theoretical/conceptual framework, key variables, current evidence, gaps, and transition to methods.",
])
for row in [4, 5]:
    set_comment(doc.tables[6], row, doc.tables[6].rows[3].cells[-1].text.split("\n"))

set_comment(doc.tables[7], 1, [
    "Databases/search engines: The section identifies peer-reviewed literature, government reports, and scientific/public health sources used to locate relevant evidence.",
    "Search terms: Search terms include HPAI H5N1, avian influenza, rice fields, wetlands, Snow Goose, waterfowl, Pacific Flyway, poultry, climate, surveillance, and One Health.",
    "Scope: The review emphasizes recent peer-reviewed literature, current outbreak-era evidence, and relevant seminal work on rice agriculture and waterfowl habitat.",
    "Limited evidence: Where direct California HPAI/rice-field evidence was limited, the review uses adjacent literature on waterfowl movement, wetlands, poultry risk, and One Health ecology.",
])
for row in [2, 3, 4]:
    set_comment(doc.tables[7], row, doc.tables[7].rows[1].cells[-1].text.split("\n"))

set_comment(doc.tables[7], 6, [
    "Theory name: One Health is identified as a guiding theory/framework.",
    "Theory origin: The section describes One Health as interdisciplinary, linking veterinary, human, environmental, and public health perspectives.",
    "Major propositions: One Health proposes that zoonotic disease risk emerges through connected human, animal, and environmental systems.",
    "Prior application: The literature shows One Health has been applied to zoonotic and HPAI risk involving wildlife, agriculture, and public health.",
    "Rationale: One Health fits this study because HPAI risk cannot be explained by poultry, wild birds, or habitat alone.",
    "Relation to RQs: The RQs test ecological and wild-bird predictors while accounting for domestic poultry exposure.",
])
for row in [7, 8, 9, 10, 11]:
    set_comment(doc.tables[7], row, doc.tables[7].rows[6].cells[-1].text.split("\n"))

set_comment(doc.tables[8], 1, [
    "Concept/phenomenon: The eco-epidemiological framework grounds the study in spillover risk at the wildlife-agriculture interface.",
    "Seminal/current writings: The section synthesizes literature on wetlands, rice fields, migratory waterfowl, poultry exposure, climate, and HPAI spread.",
    "Key statements: Risk is conceptualized as conditional on habitat availability, bird aggregation, susceptible hosts, climate, and surveillance/detection processes.",
    "Application/benefit: The framework helps organize variables and interpret why GIS convergence may appear even when adjusted county-level regressions are null.",
])
for row in [2, 3, 4]:
    set_comment(doc.tables[8], row, doc.tables[8].rows[1].cells[-1].text.split("\n"))

set_comment(doc.tables[8], 6, [
    "Constructs/methods: The review describes studies of rice-field habitat, waterfowl movement, HPAI ecology, poultry outbreak risk, climate factors, and ecological/spatial methods.",
    "Approaches/strengths/weaknesses: Prior research uses telemetry, surveillance, risk-factor modeling, and spatial analysis; strengths include ecological detail, while weaknesses include fragmented wildlife/agriculture data.",
    "Variable rationale: Rice acreage, Snow Goose activity, poultry exposure, and climate covariates are justified by literature on surrogate wetlands, migration, host exposure, and environmental persistence.",
    "Synthesis of variables: The literature supports studying rice habitat and wild-bird aggregation together with poultry exposure rather than as isolated risk factors.",
    "Synthesis of RQs: The RQs flow from the literature gap about whether county-level ecological indicators predict poultry outbreak outcomes in California.",
])
for row in [7, 8, 9, 10]:
    set_comment(doc.tables[8], row, doc.tables[8].rows[6].cells[-1].text.split("\n"))

set_comment(doc.tables[9], 1, [
    "Major themes: HPAI risk is shaped by migratory birds, surrogate wetland habitat, environmental persistence, poultry exposure, climate, and surveillance systems.",
    "Known/unknown: Literature supports biological plausibility of wild-bird and habitat pathways, but fine-scale empirical evidence linking rice, Snow Geese, and poultry outbreaks in California remains limited.",
    "Gap filled: The study extends knowledge by integrating county-winter ecological, agricultural, surveillance, and GIS data.",
    "Transition: The summary connects the literature gap to the retrospective ecological methods described in Chapter 3.",
])
for row in [2, 3, 4]:
    set_comment(doc.tables[9], row, doc.tables[9].rows[1].cells[-1].text.split("\n"))

# Chapter 3.
set_comment(doc.tables[10], 3, [
    "Purpose restated: Chapter 3 restates that the study examined rice acreage, Lesser Snow Goose activity, and their interaction in relation to poultry HPAI outbreak incidence.",
    "Preview: The chapter previews design, methodology, population/sample, data sources, variables, analysis plan, validity threats, and ethics.",
])
set_comment(doc.tables[10], 4, doc.tables[10].rows[3].cells[-1].text.split("\n"))

set_comment(doc.tables[10], 6, [
    "Variables: IVs are rice acreage and Lesser Snow Goose activity; the interaction term is rice by Snow Goose activity; the DV is poultry HPAI outbreak incidence; covariates are poultry inventory, eBird effort, temperature, and precipitation.",
    "Design/RQ connection: A retrospective ecological county-winter design matches the RQs because the study examines population-level spatial and seasonal associations.",
    "Time/resource constraints: Archival data made the study feasible without direct field sampling, telemetry collection, or farm-level recruitment.",
    "Advance knowledge: The design supports landscape-level One Health modeling using public data and reproducible spatial/statistical workflows.",
    "Intervention: Not applicable because the study did not include a treatment or manipulation.",
])
for row in [7, 8, 9, 10]:
    set_comment(doc.tables[10], row, doc.tables[10].rows[6].cells[-1].text.split("\n"))

set_comment(doc.tables[11], 2, [
    "Target population: California county-winter observations in the Pacific Flyway agricultural-wetland interface.",
    "Population size: The final primary analytic sample included 36 county-winter observations across 12 rice-producing counties; the statewide sensitivity dataset included 174 county-winter observations across 58 counties.",
])
set_comment(doc.tables[11], 3, doc.tables[11].rows[2].cells[-1].text.split("\n"))

set_comment(doc.tables[11], 5, [
    "Sampling strategy: The study used purposive/census-based ecological sampling rather than random sampling.",
    "Sample drawn: County-winter records were included when they met ecological relevance, rice-acreage, outcome-data, wild-bird activity, and covariate-completeness criteria.",
    "Sampling frame: Primary analysis required at least 1,000 CDL rice acres in any crop year from 2022-2024 and complete winter-season data; statewide sensitivity included all counties with complete data.",
    "Power/sample size: The checklist notes the planned power-analysis discussion, but final interpretation emphasizes sparse ecological outcomes and limited power for interaction testing.",
    "Effect/alpha/power: The study uses conventional quantitative criteria but interprets findings cautiously because ecological outbreak data were sparse.",
    "Tool/source: G*Power is cited in the proposal-methods discussion; final modeling decisions were based on outcome distribution and overdispersion diagnostics.",
])
for row in [6, 7, 8, 9, 10]:
    set_comment(doc.tables[11], row, doc.tables[11].rows[5].cells[-1].text.split("\n"))

set_comment(doc.tables[12], 5, [
    "Archival procedures: Data collection involved obtaining and merging public county-level datasets from USDA/CDL, DWR, eBird, APHIS/CDFA, PRISM, and USDA Agricultural Census sources.",
    "Access procedure: Data were accessed from public data portals or downloadable files and processed into county-winter analytic datasets.",
    "Permissions: No participant permissions were needed because the data were public, archival, aggregated, and non-identifiable.",
    "Historical/legal sources: Government and scientific sources were appropriate because they are standard surveillance, land-cover, climate, agricultural, and biodiversity data sources.",
])
for row in [6, 7, 8]:
    set_comment(doc.tables[12], row, doc.tables[12].rows[5].cells[-1].text.split("\n"))

set_comment(doc.tables[13], 1, [
    "Published instruments: Not applicable; the study did not use a survey, test, scale, or participant-facing instrument.",
    "Developer/year: Not applicable for instruments; archival datasets are described by source agency/platform instead.",
    "Appropriateness: Public datasets were appropriate because the study measured county-level ecological, agricultural, surveillance, and climate variables.",
    "Permission: No instrument permission was required; public data-source access was sufficient.",
    "Reliability/validity: Reliability and validity were addressed through use of established data sources and transparent operationalization rather than psychometric statistics.",
    "Prior use/populations: These data types are commonly used in ecological, spatial, agricultural, and disease surveillance research.",
    "Researcher instruments: Not applicable; variables were constructed from archival datasets rather than a researcher-developed measurement instrument.",
    "Basis for development: Variable construction was based on literature and the approved conceptual model.",
    "Reliability plan: Reproducibility was supported by documented data sources, coding, transformations, and model diagnostics.",
    "Validity plan: Construct validity was addressed by discussing proxy limitations for flooded habitat, bird activity, and surveillance-dependent detections.",
    "Sufficiency: The archival measures were sufficient for county-level ecological modeling but limited for fine-scale causal inference.",
])
for row in range(2, 12):
    set_comment(doc.tables[13], row, doc.tables[13].rows[1].cells[-1].text.split("\n"))

set_comment(doc.tables[14], 7, [
    "Operational definition: Each variable is defined at the county-winter level, including rice acreage, Snow Goose activity, poultry outbreak counts, poultry inventory, eBird effort, wild-bird detections, and climate.",
    "Measurement: Variables are measured from CDL/DWR, eBird, APHIS/CDFA, PRISM, and USDA Agricultural Census data sources.",
    "Calculation/score: Variables are summarized, transformed, and standardized as appropriate; model coefficients are interpreted as adjusted rate ratios or odds ratios depending on model type.",
])
for row in [8, 9]:
    set_comment(doc.tables[14], row, doc.tables[14].rows[7].cells[-1].text.split("\n"))

set_comment(doc.tables[14], 11, [
    "Software: R and GIS/QGIS-compatible spatial outputs were used for data processing, analysis, and visualization.",
    "Cleaning/screening: Records were cleaned, merged by county-winter, checked for completeness, transformed, and screened for sparse outcomes and overdispersion.",
    "RQs/hypotheses: The analysis plan restates RQ1-RQ3 and the null/alternative hypotheses for rice acreage, Snow Goose activity, and their interaction.",
    "Analysis plan: Negative binomial regression was used for count outcomes; logistic regression was retained as a secondary binary-outcome analysis; GIS overlays supported spatial interpretation.",
    "Statistical tests: The primary tests were adjusted negative binomial models with rate ratios, confidence intervals, and p values.",
    "Multiple tests: Results were interpreted by research question with sensitivity analyses rather than treated as independent exploratory claims.",
    "Covariates: Poultry inventory, eBird effort, temperature, and precipitation were included to reduce confounding and align with the conceptual model.",
    "Interpretation: Results were interpreted using p values, confidence intervals, rate ratios/odds ratios, model diagnostics, and spatial convergence maps.",
])
for row in [12, 13, 14, 15, 16, 17, 18]:
    set_comment(doc.tables[14], row, doc.tables[14].rows[11].cells[-1].text.split("\n"))

set_comment(doc.tables[14], 20, [
    "External validity: Findings may not generalize beyond California or similar agricultural-wetland and migratory waterfowl systems.",
    "Internal validity: Causal inference is limited because the study is nonexperimental and ecological.",
    "Construct/statistical validity: Threats include proxy measures, county-level aggregation, sparse outcomes, overdispersion, surveillance bias, and limited interaction power.",
])
for row in [21, 22]:
    set_comment(doc.tables[14], row, doc.tables[14].rows[20].cells[-1].text.split("\n"))

set_comment(doc.tables[15], 1, [
    "Data agreements: No participant or private-data agreements were required because all data were public, archival, and aggregated.",
    "Human participants: Not applicable; the study did not recruit or interact with human participants.",
    "IRB/permissions: The archival, non-identifiable design supports exemption/low-risk review; permissions were limited to public data access.",
    "Recruitment concerns: Not applicable because no recruitment occurred.",
    "Data collection concerns: Ethical concerns were limited to responsible data handling, accurate reporting, and avoiding overstatement of causal claims.",
])
for row in [2, 3, 4, 5]:
    set_comment(doc.tables[15], row, doc.tables[15].rows[1].cells[-1].text.split("\n"))

set_comment(doc.tables[15], 6, [
    "Treatment of data: Data were archival, aggregated, and stored/processed as county-level analytic files.",
    "Anonymous/confidential: No individual people, farms, or private records were identified; county-level public data were used.",
    "Protections: Files were stored locally and reported in aggregate, with no confidential participant data.",
    "Other ethics: No workplace, power, recruitment, or intervention ethical issues applied.",
])
for row in [7, 8, 9]:
    set_comment(doc.tables[15], row, doc.tables[15].rows[6].cells[-1].text.split("\n"))

set_comment(doc.tables[15], 11, [
    "Summary of method: Chapter 3 summarizes the retrospective ecological county-winter design, archival data sources, variables, sampling frame, analysis plan, validity threats, and ethics.",
    "Transition: The chapter transitions to Chapter 4 by indicating that the next chapter reports descriptive, statistical, and GIS results.",
])
set_comment(doc.tables[15], 12, doc.tables[15].rows[11].cells[-1].text.split("\n"))

# Chapter 4.
set_comment(doc.tables[17], 3, [
    "Purpose/RQs/hypotheses: Chapter 4 restates that the study examined rice acreage, Lesser Snow Goose activity, and their interaction in relation to poultry HPAI outbreak incidence.",
    "Organization preview: The chapter is organized by introduction, pilot-study applicability, data collection, descriptive results, assumption checks, RQ1-RQ3 model results, GIS convergence analysis, rice-acreage validation, and summary.",
])
set_comment(doc.tables[17], 4, doc.tables[17].rows[3].cells[-1].text.split("\n"))

set_comment(doc.tables[17], 6, [
    "Pilot results: No pilot study was conducted because the study used archival county-level secondary data.",
    "Pilot impact: No pilot study affected instrumentation, data collection, or analysis because no recruitment, intervention, or primary data collection occurred.",
])
set_comment(doc.tables[17], 7, doc.tables[17].rows[6].cells[-1].text.split("\n"))

set_comment(doc.tables[17], 9, [
    "Time frame/recruitment: Data collection used complete winter seasons from 2022-2023 through 2024-2025; no human recruitment or response rate applied.",
    "Discrepancies: The final analysis used complete winters and excluded the incomplete 2025-2026 boundary season; CDL was primary for rice acreage, with DWR used for validation.",
    "Baseline characteristics: The primary sample included 36 county-winter observations across 12 rice-producing counties; statewide sensitivity included 174 county-winter observations across all 58 counties.",
    "Representativeness: The primary sample represents rice-producing counties relevant to the study mechanism, while statewide sensitivity tests broader county-level robustness.",
    "Univariate/covariate support: Descriptive summaries and correlation screening supported inclusion of poultry inventory, eBird effort, and climate covariates.",
])
for row in [10, 11, 12, 13]:
    set_comment(doc.tables[17], row, doc.tables[17].rows[9].cells[-1].text.split("\n"))

set_comment(doc.tables[18], 1, [
    "Treatment fidelity: Not applicable because the study was retrospective, ecological, archival, and non-interventional.",
    "Adverse events: Not applicable because no treatment, intervention, participant activity, or manipulation occurred.",
])
set_comment(doc.tables[18], 2, doc.tables[18].rows[1].cells[-1].text.split("\n"))

set_comment(doc.tables[18], 4, [
    "Descriptive statistics: Chapter 4 reports outbreak counts, rice acreage, Snow Goose activity, poultry inventory, climate covariates, eBird effort, wild-bird detections, and winter trends.",
    "Assumptions: Assumption checks showed sparse and overdispersed count outcomes, supporting negative binomial regression rather than ordinary Poisson models.",
    "Findings by RQ: RQ1 rice acreage, RQ2 Snow Goose activity, and RQ3 interaction terms were not statistically significant in adjusted models.",
    "Exact statistics/p values: Chapter 4 reports rate ratios, confidence intervals, and p values for the focal predictors and sensitivity models.",
    "Confidence intervals: Confidence intervals are reported for the regression estimates and used to interpret uncertainty around the null findings.",
    "Effect sizes: Rate ratios/odds ratios serve as effect-size measures for count and binary models.",
    "Post-hoc analyses: No separate post-hoc hypothesis tests were emphasized; supplemental and sensitivity analyses were interpreted cautiously.",
    "Additional tests: GIS convergence, statewide sensitivity models, and DWR/CDL validation supported interpretation without replacing the main null tests.",
    "Tables/figures: Tables and figures summarize descriptive statistics, model results, GIS convergence, winter progression, and rice-acreage validation.",
])
for row in range(5, 13):
    set_comment(doc.tables[18], row, doc.tables[18].rows[4].cells[-1].text.split("\n"))

set_comment(doc.tables[18], 14, [
    "Answers to RQs: The null hypotheses for RQ1, RQ2, and RQ3 were not rejected in the adjusted models.",
    "Transition: Chapter 4 transitions to Chapter 5 by explaining that GIS convergence and null regression findings require cautious interpretation through the One Health framework.",
])
set_comment(doc.tables[18], 15, doc.tables[18].rows[14].cells[-1].text.split("\n"))

# Chapter 5.
set_comment(doc.tables[18], 19, [
    "Purpose/nature: Chapter 5 reiterates that the study was a quantitative retrospective ecological analysis of rice acreage, Snow Goose activity, and poultry HPAI outbreak incidence.",
    "Key findings: The chapter summarizes null adjusted findings for rice acreage, Snow Goose activity, and their interaction; increasing outbreak burden; poultry inventory as the strongest measurable predictor; and GIS convergence in selected counties.",
])
set_comment(doc.tables[18], 20, doc.tables[18].rows[19].cells[-1].text.split("\n"))

set_comment(doc.tables[18], 22, [
    "Confirm/disconfirm/extend: Findings disconfirmed statistically detectable adjusted associations for the focal predictors but extended knowledge by showing where ecological and agricultural indicators spatially converge.",
    "Framework interpretation: Findings are interpreted through One Health and disease ecology, emphasizing that risk likely depends on linked habitat, wild-bird movement, poultry exposure, climate, surveillance, and fine-scale contact pathways.",
])
set_comment(doc.tables[18], 23, doc.tables[18].rows[22].cells[-1].text.split("\n"))

set_comment(doc.tables[18], 25, [
    "Limitations: Limitations include county-level aggregation, rice acreage as an imperfect proxy for winter flooding, eBird observer effort, static poultry inventory, sparse outcomes, surveillance/reporting bias, and limited causal inference.",
])

set_comment(doc.tables[18], 27, [
    "Further research: Recommendations include finer-scale farm-buffer or landscape-cell analyses, direct inundation data, broader wild-bird/telemetry indicators, lagged seasonal models, poultry operation characteristics, and biosecurity data.",
])

set_comment(doc.tables[19], 1, [
    "Positive social change impact: The study may support targeted One Health surveillance, preparedness, cross-sector coordination, agricultural planning, food-system protection, and wildlife-poultry risk reduction.",
    "Boundaries: Implications do not claim causality or universal generalizability; they remain bounded to county-level ecological evidence and similar agricultural-wetland landscapes.",
    "Methodological/theoretical/empirical implications: The study demonstrates both the usefulness and limits of public archival county-level modeling and supports finer-scale integrated wildlife, habitat, poultry, and climate data.",
    "Practice recommendations: Practice implications include targeted surveillance in high-convergence areas, continued communication among wildlife/agriculture/public health partners, and efforts to reduce wild-bird contact with poultry.",
])
for row in [2, 3, 4, 5]:
    set_comment(doc.tables[19], row, doc.tables[19].rows[1].cells[-1].text.split("\n"))

set_comment(doc.tables[19], 7, [
    "Take-home message: The main regression findings were null, but the ecological convergence theory remains plausible; county-level data may be too coarse to detect fine-scale wildlife-poultry transmission pathways.",
])

# Final self-check / APA style.
set_comment(doc.tables[20], 4, [
    "Past tense: Chapters 1-3 were reviewed for completed-study wording where appropriate, especially Chapter 3 methods and final analytic descriptions.",
    "Method updates: Chapter 3 reflects the final retrospective ecological county-winter design, R/GIS workflow, negative binomial count models, logistic sensitivity models, and public archival data sources.",
])
set_comment(doc.tables[20], 5, doc.tables[20].rows[4].cells[-1].text.split("\n"))

set_comment(doc.tables[20], 7, [
    "Revised Study Design section: N/A. The revision clarified framing and sample language but did not introduce a major research design change requiring a new Chapter 3 section or IRB change.",
    "Data collection changes: N/A for major-change reporting; final data collection remained archival and public.",
    "Population changes: N/A for major-change reporting; the checklist notes the final primary rice-producing county sample and statewide sensitivity dataset.",
    "Analysis changes: N/A for major-change reporting; final analysis is documented as negative binomial count modeling, logistic sensitivity modeling, and GIS convergence analysis.",
])
for row in [8, 9, 10]:
    set_comment(doc.tables[20], row, doc.tables[20].rows[7].cells[-1].text.split("\n"))

set_comment(doc.tables[20], 12, [
    "Abstract consistency: The abstract includes the final method, main null findings, GIS convergence interpretation, and bounded implications.",
    "Chapter 1 consistency: N/A for a required major-change paragraph; the revision clarified framing without introducing a major design change.",
])
set_comment(doc.tables[20], 13, doc.tables[20].rows[12].cells[-1].text.split("\n"))

set_comment(doc.tables[21], 4, [
    "Citation/reference crosscheck: Citations were checked structurally against the existing bibliography, and Chapter 5 citations were limited to sources already present in the dissertation reference list.",
    "APA source formatting: References remain in APA-style author-date format and the sample template reference/appendix instruction pages were removed from the final dissertation draft.",
])
set_comment(doc.tables[21], 5, doc.tables[21].rows[4].cells[-1].text.split("\n"))

set_comment(doc.tables[21], 7, [
    "Grammar/spelling/syntax: The revised sections were reviewed for coherence, null-finding interpretation, and alignment across Chapters 1, 3, 4, and 5.",
    "Verb tense: Final-study sections were reviewed for completed-study wording, especially in Chapter 4 and Chapter 5.",
])
set_comment(doc.tables[21], 8, doc.tables[21].rows[7].cells[-1].text.split("\n"))

set_comment(doc.tables[21], 10, [
    "Headings: The final draft uses the dissertation template heading structure for Chapter 4, Chapter 5, and references.",
])

set_comment(doc.tables[21], 12, [
    "Template use: The dissertation template structure was retained, the final draft was double-spaced, and leftover template sample instruction pages were removed after the real reference list.",
])

doc.save(OUT)
print(OUT)
