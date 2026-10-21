from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt


SOURCE = Path("data/manuscripts/Quantitative Dissertation Checklist.docx")
OUT = Path("analysis_outputs/final_revised_draft/Messinger_Quantitative_Checklist_Final_Revised_Draft.docx")


def set_cell(cell, text):
    cell.text = text
    for para in cell.paragraphs:
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        for run in para.runs:
            run.font.size = Pt(9)


def update_rows(table, rows, page, comment):
    for row_idx in rows:
        row = table.rows[row_idx]
        if len(row.cells) >= 3:
            set_cell(row.cells[-2], page)
            set_cell(row.cells[-1], comment)


doc = Document(SOURCE)

front_summary = (
    "Type of Relationship: Correlational/ecological, assessed through county-winter associations and GIS spatial convergence. "
    "Key Variables: IVs: rice acreage as a proxy for potential flooded rice habitat and Lesser Snow Goose activity as a wild-bird migration indicator; "
    "DV: poultry HPAI H5N1 outbreak incidence; covariates: poultry inventory, eBird checklist effort, winter temperature, and precipitation. "
    "Population: California county-winter observations in the Pacific Flyway agricultural-wetland interface, with a primary rice-producing county sample and statewide sensitivity analysis."
)
abstract_note = (
    "Abstract summarizes the problem of HPAI H5N1 risk at the wild-bird, rice-agriculture, and poultry interface; identifies the quantitative ecological purpose; "
    "states the One Health/disease ecology framing; summarizes RQ1-RQ3; and reports the main null findings with GIS convergence, limitations, and positive social change implications."
)

for row_idx in [3, 4, 5]:
    set_cell(doc.tables[0].rows[row_idx].cells[1], front_summary)
for row_idx in range(7, 14):
    set_cell(doc.tables[0].rows[row_idx].cells[1], abstract_note)

# Chapter 1 page updates and light reframing.
ch1_intro = (
    "The introduction presents HPAI H5N1 as an agricultural, wildlife, and public health concern in California. "
    "It explains why the wild-bird migration/rice-field habitat interface matters and connects the study to One Health surveillance and positive social change."
)
update_rows(doc.tables[1], [4, 5], "1-2", ch1_intro)

background_note = (
    "The background reviews HPAI H5N1 emergence, Central Valley rice agriculture as surrogate wetland habitat, waterfowl congregation, climate variability, and the gap in integrated ecological analyses of wild-bird and agricultural spillover risk."
)
update_rows(doc.tables[2], [1], "2-5", background_note)
update_rows(doc.tables[2], [2], "4-5", background_note)
update_rows(doc.tables[2], [3], "5", background_note)

problem_note = (
    "The problem statement identifies the limited empirical evidence quantifying relationships among rice-field habitat, migratory waterfowl activity, and HPAI outbreak patterns in California, despite increasing recognition of One Health risks at the wildlife-agriculture interface."
)
update_rows(doc.tables[2], [5, 6, 7, 8], "5-6", problem_note)

purpose_note = (
    "The purpose states that this quantitative ecological study examined rice acreage, Lesser Snow Goose activity, and their interaction in relation to poultry HPAI H5N1 outbreak incidence, with poultry inventory treated as the susceptible-host covariate."
)
update_rows(doc.tables[2], [10, 11], "6-7", purpose_note)

rq_note = (
    "RQ1 tested rice acreage and poultry outbreak incidence; RQ2 tested Lesser Snow Goose activity as the wild-bird migration indicator; RQ3 tested the rice acreage by Snow Goose interaction. Null and alternative hypotheses identify the ecological predictors and poultry outbreak outcome."
)
update_rows(doc.tables[2], [13, 14], "7-9", rq_note)

theory_note = (
    "The study is grounded in One Health and eco-epidemiological/disease ecology frameworks. These frameworks connect environmental habitat, wild-bird movement, domestic poultry exposure, and zoonotic disease outcomes."
)
update_rows(doc.tables[3], [2, 3, 4], "9-10", theory_note)
update_rows(doc.tables[3], [6, 7, 8, 9], "10-11", theory_note)
nature_note = (
    "The nature of the study identifies a retrospective ecological design using county-winter archival data. It briefly lists rice acreage and Lesser Snow Goose activity as ecological predictors, poultry outbreak incidence as the outcome, and poultry/climate/effort covariates."
)
update_rows(doc.tables[3], [11, 12, 13], "11-13", nature_note)

definitions_note = (
    "Definitions clarify rice acreage/flooded rice habitat, Lesser Snow Goose activity, HPAI outbreak incidence, poultry exposure, and related ecological terms, with operational detail expanded in Chapter 3."
)
update_rows(doc.tables[4], [1, 2, 3], "13", definitions_note)
assumptions_note = (
    "Assumptions address the accuracy and comparability of archival USDA, eBird, PRISM, and surveillance datasets and the use of county-level variables as ecological proxies for wild-bird and agricultural exposure pathways."
)
update_rows(doc.tables[4], [5, 6], "13-15", assumptions_note)
scope_note = (
    "Scope and delimitations define the focus on California county-winter observations in the Central Valley/Pacific Flyway agricultural-wetland interface, with emphasis on rice-producing counties and public archival data."
)
update_rows(doc.tables[4], [8, 9, 10], "15-16", scope_note)

limitations_ch1 = (
    "Limitations include the ecological county-level design, proxy measurement of flooded rice habitat, eBird observer effort, surveillance dependence in wild-bird detections, and limited generalizability beyond similar agricultural-wetland landscapes."
)
update_rows(doc.tables[5], [1, 2, 3], "16-18", limitations_ch1)
significance_note = (
    "Significance is framed around advancing One Health surveillance by integrating agricultural habitat, wild-bird activity, poultry exposure, climate, and outbreak data to support prevention-focused planning."
)
update_rows(doc.tables[5], [5, 6, 7], "18-19", significance_note)
summary_ch1 = (
    "Chapter 1 summarizes the problem, purpose, research questions, conceptual frameworks, scope, limitations, and significance, then transitions to the literature review."
)
update_rows(doc.tables[5], [9, 10], "19-20", summary_ch1)

# Chapter 2 page updates.
ch2_intro = (
    "Chapter 2 restates the problem and purpose and establishes the relevance of HPAI H5N1 ecology, rice-field habitat, wild-bird migration, poultry exposure, and One Health surveillance."
)
update_rows(doc.tables[6], [3, 4, 5], "21-23", ch2_intro)
lit_search = (
    "The literature search strategy describes peer-reviewed and government/scientific sources, search terms related to HPAI, rice agriculture, wetlands, waterfowl, Snow Geese, poultry, climate, and One Health, and the treatment of limited direct California-specific evidence."
)
update_rows(doc.tables[7], [1, 2, 3, 4], "23-25", lit_search)
theory_ch2 = (
    "One Health and disease ecology are presented as the main frameworks. The review explains how these frameworks have been used to interpret zoonotic risk across wildlife, agricultural, environmental, and public health systems."
)
update_rows(doc.tables[7], [6, 7, 8, 9, 10, 11], "25-28", theory_ch2)
concept_ch2 = (
    "The conceptual framework links surrogate wetland habitat, wild-bird aggregation, poultry exposure, climate conditions, and HPAI outbreak occurrence. This supports the selection of rice acreage, Lesser Snow Goose activity, and poultry covariates."
)
update_rows(doc.tables[8], [1], "28", concept_ch2)
update_rows(doc.tables[8], [2, 3, 4], "28-32", concept_ch2)
lit_vars = (
    "The review synthesizes literature on HPAI ecology, migratory waterfowl, California rice fields, environmental persistence, climate variability, poultry biosecurity, and surveillance gaps. These studies justify the ecological variables and retrospective county-winter design."
)
update_rows(doc.tables[8], [6, 7, 8, 9, 10], "32-45", lit_vars)
ch2_summary = (
    "The chapter concludes that prior literature supports the plausibility of the wild-bird/rice-field exposure pathway but shows a gap in integrated county-level analyses linking ecological predictors to HPAI outbreak outcomes in California."
)
update_rows(doc.tables[9], [1], "46", ch2_summary)
update_rows(doc.tables[9], [2, 3, 4], "46-49", ch2_summary)

# Chapter 3 page updates.
ch3_intro = (
    "Chapter 3 restates the study purpose and previews the retrospective ecological design, county-winter unit of analysis, archival data sources, variables, analytic plan, validity threats, and ethical procedures."
)
update_rows(doc.tables[10], [3, 4], "50-51", ch3_intro)
design_note = (
    "The design section identifies rice acreage and Lesser Snow Goose activity as ecological predictors, poultry HPAI outbreak incidence as the outcome, and poultry inventory, eBird effort, and climate as covariates. The retrospective ecological design fits the RQs and available archival data."
)
update_rows(doc.tables[10], [6], "51", design_note)
update_rows(doc.tables[10], [7, 8, 9, 10], "51-53", design_note)
population_note = (
    "Target population: California county-winter observations within the Pacific Flyway agricultural-wetland interface. The primary analysis focuses on rice-producing counties meeting the CDL rice-acreage threshold, with a statewide sensitivity analysis."
)
update_rows(doc.tables[11], [2, 3], "55-57", population_note)
sampling_note = (
    "Sampling is purposive/census-based at the county-winter level. Inclusion criteria are ecological relevance, rice acreage threshold for the primary sample, available outbreak outcome data, migratory waterfowl activity, and complete covariates."
)
update_rows(doc.tables[11], [5, 6, 7, 8, 9, 10], "57-59", sampling_note)
archival_note = (
    "No human participants were recruited. Data were archival and publicly available from USDA/CDL, DWR, eBird, APHIS/CDFA, PRISM, and USDA Agricultural Census sources; permissions were not required beyond use of public data."
)
update_rows(doc.tables[12], [5, 6, 7, 8], "59-63", archival_note)
instr_note = (
    "No published survey or researcher-developed instrument was used. The study used archival datasets operationalized into county-winter measures; sufficiency is based on publicly documented data sources and reproducible variable construction."
)
update_rows(doc.tables[13], list(range(1, 12)), "64-68", instr_note)
oper_note = (
    "Operationalization describes how rice acreage, Lesser Snow Goose activity, poultry outbreak counts, poultry inventory, eBird checklist effort, climate covariates, and supplemental wild-bird detections were measured and transformed for analysis."
)
update_rows(doc.tables[14], [7, 8, 9], "64-68", oper_note)
analysis_note = (
    "R and GIS/QGIS-compatible tools were used for data processing, modeling, and spatial outputs. Negative binomial regression was the primary count model, logistic regression was secondary, and GIS overlays described spatial convergence."
)
update_rows(doc.tables[14], [11, 12], "68", analysis_note)
update_rows(doc.tables[14], [13, 14, 15, 16, 17, 18], "68-72", analysis_note)
validity_note = (
    "Validity threats include ecological fallacy, county-level aggregation, proxy measures for flooded habitat and bird activity, sparse outbreak outcomes, surveillance/reporting bias, and limited generalizability. Model diagnostics and sensitivity analyses addressed these issues where possible."
)
update_rows(doc.tables[14], [20, 21, 22], "73-75", validity_note)
ethics_note = (
    "The study used de-identified, public, archival county-level data and did not involve human participants. Data handling focused on reproducibility, secure storage, and transparent reporting."
)
update_rows(doc.tables[15], [1, 2, 3, 4, 5], "75", ethics_note)
update_rows(doc.tables[15], [6, 7, 8, 9], "75-76", ethics_note)
summary_ch3 = (
    "Chapter 3 summarizes the retrospective ecological design, county-winter sample, archival data sources, operational definitions, analysis plan, validity threats, and ethical procedures, then transitions to Chapter 4 results."
)
update_rows(doc.tables[15], [11, 12], "76-77", summary_ch3)

# Chapter 4.
ch4_intro = (
    "Chapter 4 briefly restates the purpose and RQ1-RQ3, clarifying that rice acreage and Lesser Snow Goose activity represent ecological/wild-bird exposure pathways and poultry outbreak incidence is the measurable outcome. The chapter previews data collection, descriptive findings, model results, GIS outputs, and summary."
)
update_rows(doc.tables[17], [3, 4], "78", ch4_intro)
pilot_note = (
    "No pilot study was conducted because the study used archival county-level secondary data, no recruitment, and no intervention. Therefore, there were no pilot results or pilot-related changes to the main analysis."
)
update_rows(doc.tables[17], [6, 7], "79", pilot_note)
data_collection_note = (
    "Data collection used complete winter seasons from 2022-2023 through 2024-2025. The primary sample included 36 county-winter observations across 12 rice-producing counties; the statewide sensitivity dataset included 174 county-winter observations across 58 counties. No human recruitment or response rate applied."
)
update_rows(doc.tables[17], [9, 10, 11, 12, 13], "79-81", data_collection_note)

treatment_note = "Not applicable. The study was retrospective, ecological, and archival, with no treatment, intervention, manipulation, or adverse events."
update_rows(doc.tables[18], [1, 2], "N/A", treatment_note)
results_note = (
    "Chapter 4 reports descriptive statistics, overdispersion diagnostics, negative binomial and logistic model results, and GIS figures. Poultry outbreaks increased over time; RQ1, RQ2, and RQ3 focal predictors were not statistically significant; poultry inventory was the most consistent predictor; GIS showed spatial convergence in selected counties."
)
update_rows(doc.tables[18], [4], "81-83", results_note)
update_rows(doc.tables[18], [5], "84-85", results_note)
update_rows(doc.tables[18], [6, 7, 8, 9, 10, 11], "86-88", results_note)
update_rows(doc.tables[18], [12], "88-92", results_note)
ch4_summary = (
    "The summary states that the null hypotheses for RQ1-RQ3 were not rejected, while GIS results suggested ecological/agricultural convergence that may be too fine-scale for county-level models to detect."
)
update_rows(doc.tables[18], [14, 15], "93-94", ch4_summary)

# Chapter 5.
ch5_intro = (
    "Chapter 5 reiterates the purpose and nature of the study and summarizes the key findings: null adjusted regression findings for rice acreage, Snow Goose activity, and their interaction; increasing outbreak burden; poultry inventory as the strongest measurable predictor; and GIS convergence."
)
update_rows(doc.tables[18], [19, 20], "95-96", ch5_intro)
interpret_note = (
    "Findings are interpreted as not rejecting the null hypotheses, while still supporting the plausibility of the One Health ecological convergence theory. The discussion explains that county-level archival proxies may be too coarse to capture fine-scale wildlife-poultry transmission pathways."
)
update_rows(doc.tables[18], [22, 23], "96-103", interpret_note)
limitations_note = (
    "Limitations include county-level ecological aggregation, imperfect rice-acreage proxy for actual winter flooding, eBird observer effort, static poultry inventory, sparse outcomes, and surveillance/reporting bias in wild-bird and outbreak records."
)
update_rows(doc.tables[18], [25], "103-104", limitations_note)
recommendations_note = (
    "Recommendations call for finer-scale farm-buffer or landscape-cell analyses, direct winter inundation measures, broader wild-bird and telemetry indicators, lagged/seasonal models, and better integration of poultry operation characteristics and biosecurity data."
)
update_rows(doc.tables[18], [27], "104-106", recommendations_note)

combined_implications_note = (
    "Implications emphasize targeted One Health surveillance where poultry exposure overlaps with rice-field habitat and waterfowl activity. "
    "Social change claims are bounded to improved preparedness, surveillance coordination, agricultural planning, and wildlife-poultry risk reduction. "
    "Methodological/theoretical implications note the value and limits of county-level archival modeling, and practice recommendations support cross-sector communication and reducing wild-bird contact with poultry in high-convergence areas."
)
update_rows(doc.tables[19], [1, 2, 3, 4, 5], "106-108", combined_implications_note)
conclusion_note = (
    "The conclusion states the take-home message: the main regression findings were null, but the theory of ecological convergence remains plausible; county-level data may be too coarse, and future One Health surveillance should integrate rice-field habitat, wild-bird movement, poultry exposure, and climate data."
)
update_rows(doc.tables[19], [7], "109-110", conclusion_note)

# Final document self-check.
final_check = (
    "Final revised draft was prepared in the dissertation template, double-spaced, and cleaned so the document ends at the actual reference list. "
    "Chapters 1-3 were reviewed for completed-study wording and alignment with the final analytic approach, including R/GIS workflow, negative binomial count models, logistic sensitivity models, and public archival data sources. "
    "Page numbers in this checklist correspond to the rendered final revised draft."
)
update_rows(doc.tables[20], [4], "1-77", final_check)
update_rows(doc.tables[20], [5], "50-77", final_check)
major_change_note = "N/A. The revision clarified framing and sample language but did not introduce a major design change requiring a new Chapter 3 Revised Study Design section."
update_rows(doc.tables[20], [7, 8, 9, 10, 12, 13], "N/A", major_change_note)
for row_idx in [14, 15]:
    for cell in doc.tables[20].rows[row_idx].cells:
        set_cell(cell, "")

# APA/style check.
apa_note = (
    "Draft was checked structurally for APA 7 dissertation-template alignment, heading consistency, double spacing, citation/reference presence, and removal of sample appendix/template instruction pages. Final Word review should update the TOC and confirm any university-specific front-matter placeholders."
)
for row_idx in [4, 5, 7, 8, 10, 12]:
    row = doc.tables[21].rows[row_idx]
    set_cell(row.cells[-1], apa_note)

doc.save(OUT)
print(OUT)
