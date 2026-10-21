from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement


SOURCE = Path("analysis_outputs/final_revised_draft/Messinger_Dissertation_Final_Revised_Draft.docx")
OUT = Path("analysis_outputs/committee2_revision/Messinger_Dissertation_Committee2_Revised_Draft.docx")


def insert_paragraph_after(paragraph: Paragraph, text: str, style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style is not None:
        new_para.style = style
    new_para.text = text
    return new_para


def replace_first(doc: Document, startswith: str, replacement: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(startswith):
            paragraph.text = replacement
            return paragraph
    raise ValueError(f"Paragraph not found: {startswith}")


def replace_exact(doc: Document, exact: str, replacement: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact:
            paragraph.text = replacement
            return paragraph
    raise ValueError(f"Paragraph not found: {exact}")


doc = Document(SOURCE)

# Committee response 1: revise Chapter 3 methods sections from proposal/future
# language into completed-study language.
replace_first(
    doc,
    "To determine the minimum number of observations required",
    "To evaluate whether the final analytic sample was adequate for the planned ecological models, an a priori sample size calculation was reviewed using G*Power 3.1 (Faul et al., 2009). Because G*Power does not provide a direct option for the negative binomial county-winter count models used in the final analysis, a fixed-model multiple regression test for R² deviation from zero was used as an approximation for the planned multivariable association test. With alpha set at .05, desired power set at .80, a medium effect size (f² = .15), and three focal predictors (rice acreage, Lesser Snow Goose activity, and the rice acreage by Snow Goose interaction), the estimated minimum sample size was approximately 77 observations.",
)

sample_para = None
for paragraph in doc.paragraphs:
    if paragraph.text.strip().startswith("To evaluate whether the final analytic sample was adequate"):
        sample_para = paragraph
        break
if sample_para is None:
    raise ValueError("Revised sample size paragraph not found")
insert_paragraph_after(
    sample_para,
    "The final primary rice-producing county dataset included 36 county-winter observations across 12 counties and three complete winter seasons, which was below the approximate benchmark from the a priori calculation. The statewide sensitivity dataset included 174 county-winter observations across all 58 California counties and exceeded this benchmark. Therefore, the primary rice-producing county models, particularly the interaction model reported in Table 4, were interpreted as potentially underpowered for detecting small to moderate independent or interaction effects. This limitation was heightened by sparse outcomes: 24 of 36 primary observations (66.7%) had zero poultry outbreaks. The sample size calculation and observed sparsity support a cautious interpretation of the null findings as absence of statistically detectable evidence in the available county-level data, rather than proof that rice-field habitat or wild-bird activity has no ecological relationship to outbreak risk.",
    sample_para.style,
)

replace_first(
    doc,
    "Because this study relies exclusively on archival data",
    "Because this study relied exclusively on archival data from public, governmental, and scientific sources, no procedures involving the recruitment or participation of human subjects were required. No interviews, surveys, or direct interactions occurred, and no identifiable information was collected. Instead, the procedures focused on identifying, obtaining, validating, and integrating secondary datasets that characterized environmental conditions, wildlife activity, and zoonotic disease events across California counties (Chanda et al., 2025; Harvey et al., 2023).",
)
replace_first(
    doc,
    "Data collection begins with identifying suitable data sources",
    "Data collection began with identifying suitable data sources for each variable. Rice acreage was obtained from the USDA Cropland Data Layer, which provides annual, satellite-derived crop classification maps at high spatial resolution. These data were supplemented with crop-mapping information from the California Department of Water Resources, which monitors rice acreage and agricultural land use relevant to post-harvest flooding and migratory bird habitat (Matthews et al., 2022; Rosenberg et al., 2022; Strum et al., 2013).",
)
replace_first(
    doc,
    "Lesser Snow Goose density data will be obtained",
    "Lesser Snow Goose activity data were obtained from eBird.org, an open-access community-science platform hosted by the Cornell Lab of Ornithology, using spatial and temporal filters for species, date, and location. To further contextualize spatial patterns, eBird observations were interpreted alongside migratory telemetry and movement literature used in prior work on waterfowl movement and avian influenza risk (McDuie et al., 2022, 2024; Overton & Casazza, 2023; Teitelbaum et al., 2024).",
)
replace_first(
    doc,
    "Outbreak data for HPAI H5N1 will be obtained",
    "Outbreak data for HPAI H5N1 were obtained from the USDA Animal and Plant Health Inspection Service (APHIS) outbreak dashboard, which provides downloadable records of laboratory-confirmed avian influenza detections by county and date. Where available, these data were cross-checked against state-level reports from the California Department of Food and Agriculture (CDFA) and official press releases issued by USDA APHIS and wildlife health authorities (Green et al., 2023; Kan-Rice, 2025; Lowen et al., 2025).",
)
replace_first(
    doc,
    "Climatic covariates, including precipitation and temperature, will be obtained",
    "Climatic covariates, including precipitation and temperature, were obtained from the PRISM Climate Group, a widely used source of spatial climate datasets for ecological and epidemiological modeling (Ferenczi et al., 2021; Wu et al., 2020). Covariates related to poultry production were drawn from the USDA Agricultural Census, which provides county-level data on poultry operations by type and size, consistent with prior analyses of HPAI risk factors in commercial flocks (Green et al., 2023; Patyk et al., 2023).",
)
replace_first(
    doc,
    "All datasets will be downloaded",
    "All datasets were downloaded in CSV, Excel, or spatial data formats. After acquisition, data were cleaned and organized by county-winter season, then merged using unique identifiers based on county FIPS codes and season-year combinations. Geographic data were integrated into GIS-compatible spatial outputs to generate overlays and evaluate geographic clustering and spatial convergence among habitat, waterfowl activity, poultry exposure, and outbreak incidence (Chanda et al., 2025; Prosser et al., 2022). Statistical analyses were conducted in R to support reproducible modeling, diagnostics, and visualization.",
)
replace_first(
    doc,
    "No special permissions are required",
    "No special permissions were required to use these datasets, as all sources were publicly available, managed by government agencies, or hosted by academic institutions. Full citations and documentation for each data source were provided in the dissertation, and data-handling procedures followed best practices for reproducible, transparent ecological and epidemiological research (Finucane et al., 2014; Zinsstag et al., 2011).",
)
replace_exact(doc, "The following archival data sources will be used:", "The following archival data sources were used:")
replace_first(
    doc,
    "All data are publicly accessible",
    "All data were publicly accessible, and no individual-level or identifiable records were included. Datasets were stored on a secure, encrypted drive during analysis, and all sources were cited in accordance with APA 7 guidelines. Because the research involved no intervention, manipulation, or access to private data, the study was reviewed under the university's ethical research procedures for archival, non-human-subjects research.",
)
replace_first(
    doc,
    "Proportion of flooded rice cropland.",
    "Proportion of flooded rice cropland.\nThis variable captured the extent of rice acreage used as a proxy for potential post-harvest flooded rice habitat in each county-season. Rice acreage was extracted from the USDA Cropland Data Layer, which uses remote sensing and satellite imagery to classify crop types annually. Seasonal crop-mapping information was supplemented with data from the California Department of Water Resources (Matthews et al., 2022; Rosenberg et al., 2022; Strum et al., 2013). The variable was entered as a continuous predictor, with transformations and standardization applied where appropriate for regression modeling.",
)
replace_first(
    doc,
    "To address Research Question 3, an interaction term will be created",
    "To address Research Question 3, an interaction term was created by multiplying the centered values of rice acreage and Lesser Snow Goose activity. This term was used to assess whether the association between rice-field habitat and HPAI H5N1 outbreak incidence varied with winter goose activity, consistent with prior eco-epidemiological models that incorporate interaction effects between habitat and host movement (Chanda et al., 2025; McDuie et al., 2024; Teitelbaum et al., 2024).",
)
replace_exact(doc, "Two sets of covariates will be included to control for potential confounding:", "Two sets of covariates were included to control for potential confounding:")
replace_first(
    doc,
    "Climate anomalies.",
    "Climate anomalies.\nSeasonal precipitation and temperature data were obtained from the PRISM Climate Group and summarized for each county-winter season. These variables captured unusual warm, cold, wet, or dry conditions that may influence both habitat availability and virus persistence (Choudhary et al., 2024; Ferenczi et al., 2021; Wu et al., 2020). They were incorporated as continuous covariates.",
)
replace_first(
    doc,
    "Lesser Snow Goose activity.",
    "Lesser Snow Goose activity.\nThis variable represented the relative winter activity of Lesser Snow Geese (Anser caerulescens) in each county-winter season and served as the primary wild-bird migration indicator for RQ2 and RQ3. Goose observations were obtained from eBird.org and filtered by species, date (November-February), and county. To account for differences in observer effort across counties, activity was standardized as the mean number of Lesser Snow Geese reported per complete eBird checklist per county-winter season (Baldwin et al., 2022; Carter et al., 2025). This measure provided a relative index of species presence and aggregation while partially controlling for variation in reporting effort. Spatial and temporal patterns were interpreted alongside telemetry and movement literature documenting major staging and wintering areas (McDuie et al., 2022, 2024; Overton & Casazza, 2023; Sullivan et al., 2025).",
)
replace_first(
    doc,
    "All data will be coded in SPSS or R",
    "All data were coded in R, and unique identifiers were used to merge datasets by county and season (FIPS codes and season–year combinations). A detailed data dictionary documented the operational definitions, units, transformations, and coding schemes for each variable to maintain consistency and reproducibility across all analytical steps (Finucane et al., 2014; Zinsstag et al., 2011).",
)
replace_first(
    doc,
    "All data analyses will be conducted",
    "All data analyses were conducted in R. R was used for data cleaning, descriptive statistics, regression modeling, diagnostics, and visualization because it supported reproducible negative binomial regression, logistic sensitivity analyses, and spatial-data workflows. GIS processing and visualization were completed using GIS-compatible spatial outputs rather than ArcGIS Pro, which kept the analysis reproducible without requiring proprietary software.",
)
replace_first(
    doc,
    "Raw data from public sources",
    "Raw data from public sources such as USDA, eBird, APHIS, DWR, and PRISM were cleaned and merged using county-season identifiers. Records were checked for missing values, duplicates, and format inconsistencies. Variables were evaluated for sparse outcomes, outliers, distributional shape, overdispersion, and multicollinearity. Counties or county-winter observations with incomplete boundary-season data were excluded from the final primary analysis, and transformations were documented in the analysis files.",
)
replace_first(
    doc,
    "Because the dependent variable, HPAI H5N1 outbreak occurrence",
    "Because the dependent variable, HPAI H5N1 poultry outbreak incidence, was a count variable measured as the number of outbreaks per county-winter season, Poisson regression was screened as an initial count-modeling approach. The data showed overdispersion, so negative binomial regression was used for the primary count models. Logistic regression was retained as a secondary sensitivity analysis for the binary outcome of whether any poultry outbreak occurred.",
)
replace_exact(doc, "The model will include:", "The adjusted models included:")
replace_exact(doc, "Flooded rice acreage (continuous)", "Rice acreage, log-transformed and standardized where appropriate")
replace_exact(doc, "Goose density (continuous)", "Lesser Snow Goose activity, log-transformed and standardized where appropriate")
replace_exact(doc, "An interaction term (rice × goose, using centered variables)", "The rice acreage by Lesser Snow Goose activity interaction term")
replace_exact(doc, "Covariates (see below)", "Poultry inventory, eBird checklist effort, winter temperature, and winter precipitation covariates")
replace_first(
    doc,
    "Model fit will be evaluated",
    "Model fit was evaluated using standard criteria, including AIC, BIC, log-likelihood values, overdispersion screening, convergence diagnostics, and residual spatial autocorrelation diagnostics where applicable. Collinearity diagnostics were also conducted to confirm that predictors could be interpreted without severe multicollinearity.",
)
replace_first(
    doc,
    "If multiple models are run",
    "No post hoc family of exploratory hypothesis tests was used to replace the a priori research-question tests. Statewide, logistic, GIS, and DWR/CDL comparison analyses were interpreted as sensitivity or supporting analyses rather than as new primary hypotheses. P values were therefore interpreted cautiously and in relation to the prespecified research questions, model diagnostics, and confidence intervals.",
)
replace_first(
    doc,
    "Several potential covariates will be evaluated",
    "Covariates were selected based on the theoretical and empirical literature, data availability, and the need to address plausible confounding. These included:",
)
replace_exact(doc, "Poultry farm density (as a proxy for susceptible hosts)", "Poultry inventory as a proxy for susceptible domestic host exposure")
replace_exact(doc, "Temperature and precipitation anomalies (as environmental factors that may alter virus viability or bird migration)", "Winter temperature and precipitation as environmental conditions that may alter habitat availability, virus viability, or bird movement")
replace_exact(doc, "Wetland proximity or surface water coverage (as landscape-level risk factors)", "eBird checklist effort to reduce bias from unequal wildlife observation intensity")
replace_first(
    doc,
    "Covariates will be included if prior studies",
    "Covariates were retained when prior studies identified them as relevant or when they were necessary to distinguish the focal ecological exposures from domestic poultry exposure, observer effort, and winter climate conditions.",
)
replace_first(
    doc,
    "Model outputs will be interpreted",
    "Model outputs were interpreted using coefficients, rate ratios or odds ratios, 95% confidence intervals, and p values. A significance level of alpha = .05 was used for hypothesis testing. Interaction results were interpreted using the interaction rate ratio together with confidence intervals and GIS-based convergence findings. All results were presented with supporting tables and visualizations to enhance interpretation and policy relevance.",
)
replace_exact(doc, "The final regression model will take the form:", "The final count model took the general form:")
replace_exact(
    doc,
    "ln(λ) = β₀ + β₁(Rice) + β₂(Goose) + β₃(Rice × Goose) + β₄(PoultryDensity) + β₅(Climate) + ε",
    "ln(λ) = β₀ + β₁(Rice) + β₂(Goose) + β₃(Rice × Goose) + β₄(PoultryInventory) + β₅(ChecklistEffort) + β₆(Temperature) + β₇(Precipitation) + ε",
)
replace_first(
    doc,
    "Where ln(λ) is the natural log",
    "Where ln(λ) is the natural log of the expected outbreak count, and ε is the error term. The interaction term (Rice × Goose) tested whether the association between rice acreage and outbreak count varied by goose activity. Predictors were transformed, centered, or standardized where appropriate before interaction testing. Model fit was assessed using AIC, BIC, log-likelihood, overdispersion screening, and residual diagnostics.",
)
replace_first(
    doc,
    "Because this study uses archival, ecological-level data",
    "Because this study used archival, ecological-level data and a non-experimental design, certain threats to validity were acknowledged and addressed. Although internal threats were minimized by the lack of experimental manipulation, threats to construct, statistical, and external validity remained and required careful consideration.",
)
replace_first(
    doc,
    "External validity concerns how well",
    "External validity concerns how well the findings apply beyond the study sample. This study focused on California counties within specific ecological and agricultural zones, such as the Central Valley and the Pacific Flyway. As a result, the findings may not apply to other states or countries with different agricultural systems, bird migration patterns, or land use practices. However, by using multiple county-winter observations across complete winter seasons, the study enhanced its temporal and spatial relevance within similar agro-ecological contexts. Conclusions were framed within the scope of California and similar migratory corridors to prevent overgeneralization.",
)
replace_first(
    doc,
    "Construct validity refers to how well",
    "Construct validity refers to how well operationalized variables reflect the theoretical concepts being studied. This issue was addressed by operationalizing constructs using validated sources: rice acreage was measured with satellite-derived USDA CDL data and DWR crop-mapping context, goose activity was derived from effort-standardized eBird observations, and outbreak data came from laboratory-confirmed government reports. However, proxies such as average birds per checklist may imperfectly reflect true species density, and rice acreage does not directly measure the timing or duration of post-harvest flooding within counties. These limitations were acknowledged in the discussion of findings.",
)
replace_first(
    doc,
    "This validity domain concerns how well",
    "This validity domain concerns how well statistical analyses correctly identified relationships among variables and avoided errors such as Type I (false positives) or Type II (false negatives). Threats included small sample size in the primary rice-producing county analysis, sparse outbreak counts, multicollinearity among predictors, and model misspecification. These threats were reduced by:",
)
replace_exact(doc, "Using logistic regression, which is suitable for binary outcomes.", "Using negative binomial regression for overdispersed count outcomes and logistic regression for binary sensitivity analyses.")
replace_exact(doc, "Testing for multicollinearity before final model estimation", "Testing for multicollinearity and overdispersion before final model estimation.")
replace_exact(doc, "Performing power analysis to ensure enough observational units over the 5-year span.", "Reviewing sample size and power constraints, then interpreting sparse-outcome models cautiously.")
replace_first(
    doc,
    "Spatial autocorrelation (non-independence of observations",
    "Spatial autocorrelation (non-independence of observations in nearby counties) was examined using spatial diagnostic tools such as Moran's I. No significant residual spatial autocorrelation was detected in the statewide sensitivity model, which supported the use of the final regression specification while preserving a cautious interpretation of county-level ecological findings.",
)
replace_first(
    doc,
    "This study analyzes secondary, archival data",
    "This study analyzed secondary, archival data and did not involve direct contact with human participants or the collection of personally identifiable information. Therefore, the ethical risks were minimal. Ethical considerations related to data use, transparency, confidentiality, and institutional review were addressed to maintain the study's integrity.",
)
replace_first(
    doc,
    "Because this research does not involve human subjects",
    "Because this research did not involve human subjects, the study qualified for review under archival, non-human-subjects research procedures consistent with federal guidelines and Walden University's ethical research policies. Public, aggregate datasets were accessed, processed, and analyzed only after the appropriate university review requirements were completed.",
)
replace_first(
    doc,
    "To ensure the research remains secure",
    "To ensure the research remained secure, all data were kept on an encrypted, password-protected device during analysis. No personal names, addresses, or identifying information were stored or disclosed. Results were shared only in aggregated forms, such as county summaries or visualizations, preventing the re-identification of individuals or specific farm operations.",
)
replace_first(
    doc,
    "Furthermore, proper citation and attribution will be maintained",
    "Furthermore, proper citation and attribution were maintained for all data sources, and all mapping products and spatial analyses were based on publicly available GIS layers. The final dissertation includes a comprehensive reference list and, where applicable, direct links to original data repositories to ensure transparency and reproducibility.",
)
replace_first(
    doc,
    "In summary, this study adheres",
    "In summary, this study adhered to ethical principles of respect, integrity, transparency, and data protection, even without human participants. All research followed Walden University's ethical standards for doctoral research.",
)
replace_first(
    doc,
    "The methodology specifies California counties",
    "The methodology specified California counties as the unit of analysis, outlined inclusion criteria, and described the use of a purposive census approach rather than random sampling. Key variables were operationalized using validated sources, including the USDA Cropland Data Layer, eBird, APHIS wild-bird and poultry outbreak surveillance, PRISM climate data, and USDA Agricultural Census poultry covariates. Negative binomial regression was used as the primary count-modeling approach, with logistic regression retained as a secondary binary-outcome analysis. Spatial overlays were used to interpret convergence among rice-field habitat, wild-bird activity, poultry exposure, and outbreak burden.",
)
replace_first(
    doc,
    "The next chapter will present",
    "Chapter 4 presents the results of the statistical and spatial analyses, including descriptive summaries, regression model outputs, and visualizations of outbreak patterns and ecological predictors.",
)

# Committee response 2: make Table 4 sample size explicit at the table.
replace_first(
    doc,
    "Note. Predictors were standardized after log transformation where appropriate. Models adjusted",
    "Note. Primary models used n = 36 county-winter observations; statewide sensitivity models used n = 174. Predictors were transformed/standardized and covariate-adjusted. RR = rate ratio; CI = confidence interval.",
)

seen_chapter_4 = False
for paragraph in doc.paragraphs:
    if paragraph.text.strip() == "Chapter 4: Results":
        seen_chapter_4 = True
    elif seen_chapter_4 and paragraph.text.strip() == "Table 4":
        paragraph.paragraph_format.page_break_before = True
        break

doc.save(OUT)
print(OUT)
